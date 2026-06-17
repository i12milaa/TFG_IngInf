"""Genera presentacion_tfg.pptx para la defensa del TFG."""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE_TYPE

DARK_BLUE  = RGBColor(0x1A, 0x3A, 0x6B)
MID_BLUE   = RGBColor(0x2E, 0x6D, 0xA4)
LIGHT_BLUE = RGBColor(0xD6, 0xE4, 0xF0)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
DARK_GRAY  = RGBColor(0x33, 0x33, 0x33)
ACCENT     = RGBColor(0x5B, 0x9B, 0xD5)

FIGURAS = r"d:\Universidad\TFG\TFG_Part6\latex\figuras"
OUTPUT  = r"d:\Universidad\TFG\TFG_Part6\presentacion_tfg.pptx"

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)

BLANK = prs.slide_layouts[6]


def new_slide():
    return prs.slides.add_slide(BLANK)


def bg(slide, color=WHITE):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def rect(slide, l, t, w, h, color, line_color=None):
    from pptx.util import Pt as _Pt
    s = slide.shapes.add_shape(
        1, Inches(l), Inches(t), Inches(w), Inches(h)
    )
    s.fill.solid()
    s.fill.fore_color.rgb = color
    if line_color:
        s.line.color.rgb = line_color
        s.line.width = _Pt(1)
    else:
        s.line.fill.background()
    return s


def tb(slide, text, l, t, w, h, size=16, bold=False, color=DARK_GRAY,
       align=PP_ALIGN.LEFT, italic=False):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return box


def bullets(slide, items, l, t, w, h, size=16, color=DARK_GRAY, indent=False):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        bullet = "    –" if item.startswith("    ") else "•"
        p.text = f"{bullet} {item.strip()}"
        p.space_before = Pt(3)
        for run in p.runs:
            run.font.size = Pt(size)
            run.font.color.rgb = color


def comparison_table(slide, headers, rows, l, t, w, h, highlight_row=None):
    n_rows = len(rows) + 1
    n_cols = len(headers)
    shape = slide.shapes.add_table(n_rows, n_cols, Inches(l), Inches(t), Inches(w), Inches(h))
    table = shape.table

    col_w = [w * 0.42, w * 0.20, w * 0.38]
    for i, cw in enumerate(col_w[:n_cols]):
        table.columns[i].width = Inches(cw)

    for c, htext in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = htext
        cell.fill.solid()
        cell.fill.fore_color.rgb = DARK_BLUE
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        for run in p.runs:
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = WHITE

    for r, row_data in enumerate(rows, start=1):
        is_hl = (highlight_row is not None and r - 1 == highlight_row)
        for c, val in enumerate(row_data):
            cell = table.cell(r, c)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = LIGHT_BLUE if is_hl else WHITE
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER if c > 0 else PP_ALIGN.LEFT
            for run in p.runs:
                run.font.size = Pt(12)
                run.font.bold = is_hl
                run.font.color.rgb = DARK_BLUE if is_hl else DARK_GRAY
    return shape


def img(slide, name, l, t, w=None, h=None):
    path = os.path.join(FIGURAS, name)
    if not os.path.exists(path):
        print(f"  [WARN] imagen no encontrada: {name}")
        return None
    kw = {}
    if w: kw["width"]  = Inches(w)
    if h: kw["height"] = Inches(h)
    return slide.shapes.add_picture(path, Inches(l), Inches(t), **kw)


def header(slide, title, subtitle=None):
    rect(slide, 0, 0, 13.33, 1.35, DARK_BLUE)
    rect(slide, 0, 1.35, 13.33, 0.06, MID_BLUE)
    tb(slide, title, 0.35, 0.1, 12.6, 0.85, size=28, bold=True,
       color=WHITE, align=PP_ALIGN.LEFT)
    if subtitle:
        tb(slide, subtitle, 0.35, 0.9, 12.6, 0.38, size=13,
           color=LIGHT_BLUE, align=PP_ALIGN.LEFT)


# ================================================================
# SLIDE 1 — PORTADA
# ================================================================
s = new_slide()
bg(s, DARK_BLUE)
rect(s, 0, 0, 13.33, 0.55, RGBColor(0x0F, 0x25, 0x4A))
rect(s, 0, 1.15, 13.33, 0.07, MID_BLUE)
rect(s, 0, 3.2, 13.33, 0.07, MID_BLUE)

tb(s, "Sistema de Vigilancia Distribuido",
   0.5, 1.35, 12.3, 0.9, size=34, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
tb(s, "Basado en Red de Radares Cooperativos",
   0.5, 2.2, 12.3, 0.8, size=28, bold=True, color=LIGHT_BLUE, align=PP_ALIGN.CENTER)

tb(s, "Trabajo de Fin de Grado · Grado en Ingeniería Informática",
   0.5, 3.4, 12.3, 0.5, size=15, color=LIGHT_BLUE, align=PP_ALIGN.CENTER)
tb(s, "Escuela Politécnica Superior de Córdoba · Universidad de Córdoba",
   0.5, 3.85, 12.3, 0.5, size=14, color=LIGHT_BLUE, align=PP_ALIGN.CENTER)

rect(s, 0, 4.6, 13.33, 0.06, MID_BLUE)

tb(s, "Alejandro Millán de Lara",
   0.5, 4.75, 7, 0.5, size=17, bold=True, color=WHITE)
tb(s, "Tutores: Dr. Fernando León García · Cristina Martínez Ruedas",
   0.5, 5.25, 10, 0.45, size=13, color=LIGHT_BLUE)
tb(s, "Junio 2026",
   10.8, 4.75, 2.3, 0.5, size=14, color=LIGHT_BLUE, align=PP_ALIGN.RIGHT)

# ================================================================
# SLIDE 2 — ÍNDICE
# ================================================================
s = new_slide()
bg(s, WHITE)
header(s, "Estructura de la presentación")

col1 = [
    "1. Motivación y contexto",
    "2. Objetivos",
    "3. El sistema de un vistazo",
    "4. Hardware del nodo",
    "5. Diseño mecánico",
    "6. Arquitectura general",
    "7. Protocolo TDMA dinámico",
]
col2 = [
    "8.  Protocolo binario nodo–servidor",
    "9.  Firmware: arquitectura interna",
    "10. Seguimiento y trilateración",
    "11. Interfaz gráfica",
    "12. Resultados experimentales",
    "13. Conclusiones y líneas futuras",
]
bullets(s, col1, 0.6, 1.65, 5.8, 5.5, size=18)
rect(s, 6.7, 1.6, 0.05, 5.6, LIGHT_BLUE)
bullets(s, col2, 7.0, 1.65, 6.0, 5.5, size=18)

# ================================================================
# SLIDE 3 — MOTIVACIÓN
# ================================================================
s = new_slide()
bg(s, WHITE)
header(s, "Motivación y contexto", "El problema de coordinar varios sensores de bajo coste")

bullets(s, [
    "Un sensor único cubre solo un sector angular limitado y no resuelve",
    "    ambigüedades de posición entre objetos a distancias similares",
    "Desplegar varios sensores soluciona la cobertura, pero genera un",
    "    problema físico nuevo: la interferencia acústica (crosstalk)",
    "El eco de un nodo puede ser captado por el receptor de otro y",
    "    confundirse con su propia reflexión → lecturas espurias",
    "No es un fallo de software: las ondas se propagan sin lógica de",
    "    control, así que la coordinación debe ser explícita",
], 0.5, 1.6, 6.1, 5.3, size=15)

comparison_table(
    s,
    headers=["Enfoque evaluado", "Crosstalk", "Viabilidad en este proyecto"],
    rows=[
        ("Separación en frecuencia", "Nulo", "No: exige transductores especiales"),
        ("Barrido secuencial", "Nulo", "No: latencia crece con N nodos"),
        ("Contención CSMA/CA", "Alto", "No: vel. del sonido invalida la escucha"),
        ("TDMA dinámico + reutiliz. espacial", "Nulo", "Sí: elegido en este proyecto"),
    ],
    l=6.85, t=1.9, w=6.0, h=3.0,
    highlight_row=3,
)

tb(s, "Fuente: comparativa de enfoques de acceso al medio acústico (cap. Antecedentes)",
   6.85, 5.1, 6.0, 0.5, size=11, italic=True, color=DARK_GRAY)

# ================================================================
# SLIDE 4 — OBJETIVOS
# ================================================================
s = new_slide()
bg(s, WHITE)
header(s, "Objetivos del trabajo")

bullets(s, [
    "Construir un prototipo físico funcional con 3 nodos radar",
    "Diseñar un protocolo TDMA dinámico con reutilización espacial",
    "Implementar un servidor central de orquestación en Python",
    "Desarrollar un algoritmo de trilateración cooperativa",
    "Crear una interfaz gráfica de visualización en tiempo real",
    "Validar el sistema mediante pruebas sobre el prototipo físico real",
], 0.5, 1.55, 6.6, 5.5, size=17)

img(s, "Diseño_General_Funcionamiento.png", 7.4, 1.55, w=5.7)

# ================================================================
# SLIDE 5 — EL SISTEMA DE UN VISTAZO
# ================================================================
s = new_slide()
bg(s, WHITE)
header(s, "El sistema de un vistazo")

tb(s,
   "Tres nodos radar desplegados en triángulo. Cada nodo barre su sector "
   "de forma autónoma y reporta detecciones al servidor central, que coordina "
   "los disparos y estima posiciones cooperativamente.",
   0.5, 1.55, 5.8, 1.6, size=16, color=DARK_GRAY)

img(s, "AnexoDespliegueFisico.png", 0.5, 3.2, w=5.5)
img(s, "DiagramaSistemaCompleto.png", 6.5, 1.55, w=6.6)

# ================================================================
# SLIDE 6 — HARDWARE DEL NODO
# ================================================================
s = new_slide()
bg(s, WHITE)
header(s, "Hardware del nodo radar")

components = [
    ("ESP32 dual-core", "hw_esp32.png",    "MCU WiFi, 240 MHz\n2 núcleos independientes"),
    ("HC-SR04",         "hw_hcsr04.png",   "Sensor ultrasónico\nRango: 2 – 400 cm"),
    ("NEMA 17",         "hw_nema17.png",   "Motor paso a paso\n200 pasos/vuelta"),
    ("A4988",           "hw_a4988.png",    "Driver con micropasos\nhasta 1/16 de paso"),
]

cw = 2.95
for i, (name, im, desc) in enumerate(components):
    x = 0.35 + i * (cw + 0.18)
    rect(s, x, 1.55, cw, 5.6, LIGHT_BLUE)
    img(s, im, x + 0.15, 1.7, w=cw - 0.3)
    tb(s, name, x + 0.1, 4.85, cw - 0.2, 0.5,
       size=13, bold=True, color=DARK_BLUE, align=PP_ALIGN.CENTER)
    tb(s, desc, x + 0.1, 5.35, cw - 0.2, 1.0,
       size=12, color=DARK_GRAY, align=PP_ALIGN.CENTER)

# ================================================================
# SLIDE 7 — DISEÑO MECÁNICO
# ================================================================
s = new_slide()
bg(s, WHITE)
header(s, "Diseño mecánico", "Estructura impresa en 3D")

bullets(s, [
    "Estructura diseñada en CAD e impresa en PLA",
    "Soporte del sensor sobre el eje del motor paso a paso",
    "Reed switch magnético para homing a 0°",
    "3 iteraciones de diseño hasta la versión final",
    "Coste de material: <5 € por nodo",
], 0.5, 1.55, 4.8, 3.5, size=17)

img(s, "IteracionesDiseño.png",  0.4, 4.5,  w=4.8)
img(s, "Diseño3DFinal.png",      5.5, 1.55, w=3.5)
img(s, "AnexoNodoFront.png",     9.2, 1.55, w=3.9)

# ================================================================
# SLIDE 8 — ARQUITECTURA GENERAL
# ================================================================
s = new_slide()
bg(s, WHITE)
header(s, "Arquitectura general del sistema")

img(s, "DiagramaCapas.png",        0.4, 1.55, w=6.0)
img(s, "ModeloClienteServidor.png", 6.6, 1.55, w=6.5)

tb(s, "Arquitectura en 3 capas: firmware embebido (C++) · servidor central (Python) · GUI de visualización (Tkinter)",
   0.5, 6.75, 12.3, 0.55, size=13, color=DARK_GRAY, align=PP_ALIGN.CENTER)

# ================================================================
# SLIDE 9 — PROTOCOLO TDMA
# ================================================================
s = new_slide()
bg(s, WHITE)
header(s, "Protocolo TDMA dinámico", "Solución al crosstalk sin hardware especializado")

bullets(s, [
    "El servidor calcula qué pares de nodos son geométricamente incompatibles (zona sucia)",
    "Criterio: solapamiento de conos de emisión dentro de CLEAN_LIMIT = 30°",
    "Solo los pares conflictivos reciben ranuras temporales diferenciadas",
    "El resto de nodos puede disparar en paralelo → reutilización espacial",
    "La asignación se recalcula dinámicamente ante reconexiones o cambios angulares",
], 0.5, 1.55, 6.4, 4.2, size=16)

img(s, "FuncionamientoGeneralZonas.png",  0.5, 5.0,  w=5.5)
img(s, "DiagramaFlujoOrquestacion.png",   7.0, 1.55, w=6.1)

# ================================================================
# SLIDE 10 — PROTOCOLO BINARIO
# ================================================================
s = new_slide()
bg(s, WHITE)
header(s, "Protocolo binario nodo–servidor")

bullets(s, [
    "Comunicación TCP, mensajes de longitud fija con cabecera de 3 bytes (type + length)",
    "Estructuras empaquetadas con #pragma pack → misma representación en ESP32 y Python",
    "Handshake: HELLO → ASSIGN_ID → espera inicio de supertrama",
    "Opcodes: START_SUPERFRAME, ASSIGN_SLOT, REQUEST_REPORT, REQUEST_ANGLE",
    "Timeout de servidor inactivo 5 s + contador de fallos TCP → reinicio automático",
], 0.5, 1.55, 6.4, 4.0, size=16)

img(s, "DiagramaFlujoDescubrimiento.png", 0.5, 5.0,  w=4.5)
img(s, "DiagramaUMLProtocolo.png",        6.7, 1.55, w=6.4)

# ================================================================
# SLIDE 11 — FIRMWARE INTERNO
# ================================================================
s = new_slide()
bg(s, WHITE)
header(s, "Firmware: arquitectura interna", "FreeRTOS sobre ESP32 dual-core")

bullets(s, [
    "Core 0 → TaskComms: conexión TCP y protocolo con el servidor",
    "Core 1 → TaskRadar: control determinista del motor y medición ToF por IRQ",
    "Colas FreeRTOS como canal de comunicación entre tareas:",
    "    HwCommand (Comms → Radar): HOME, MOVE, EXECUTE_SLOT, SYNC_POS",
    "    HwResult  (Radar → Comms): distancia medida, ángulo alcanzado",
    "FSM de alto nivel: CONNECTING → CONNECTED → OPERATING",
    "Perfiles por MAC address: parámetros de hardware individuales por nodo",
], 0.5, 1.55, 6.3, 5.5, size=16)

img(s, "ArquitecturaCoresNodo.png", 6.8, 1.55, w=6.3)

# ================================================================
# SLIDE 12 — SEGUIMIENTO Y TRILATERACIÓN
# ================================================================
s = new_slide()
bg(s, WHITE)
header(s, "Seguimiento y trilateración cooperativa")

bullets(s, [
    "Modo SEARCH: barrido continuo buscando nuevas detecciones",
    "Modo TRACK: enfoca el sensor en el ángulo de la última detección (TRACK_LOCK = 6 ciclos)",
    "Zona de solapamiento: ≥ 2 nodos detectan el mismo objeto simultáneamente",
    "Trilateración: intersección de circunferencias de radio = distancia medida por cada nodo",
    "Posición estimada como el punto más cercano al conjunto de circunferencias",
], 0.5, 1.55, 6.3, 4.0, size=16)

img(s, "CalculoTrilateracion.png",           0.5, 5.0,  w=4.5)
img(s, "FuncionamientoGeneralTrilateracion.png", 6.7, 1.55, w=6.4)

# ================================================================
# SLIDE 13 — GUI
# ================================================================
s = new_slide()
bg(s, WHITE)
header(s, "Interfaz gráfica (Tkinter)")

tb(s, "Modo Diseño", 0.5, 1.5, 5.9, 0.45,
   size=15, bold=True, color=DARK_BLUE, align=PP_ALIGN.CENTER)
img(s, "ModoDiseñoGUI.png", 0.5, 1.97, w=5.9)

tb(s, "Modo Telemetría", 7.0, 1.5, 6.0, 0.45,
   size=15, bold=True, color=DARK_BLUE, align=PP_ALIGN.CENTER)
img(s, "ModoTelemetriaGUI.png", 7.0, 1.97, w=6.0)

tb(s, "Configuración de geometría y umbrales DEFCON · Visualización en tiempo real · "
      "Arranque/parada del servidor · Persistencia de sesión entre ejecuciones",
   0.5, 6.65, 12.3, 0.65, size=13, color=DARK_GRAY, align=PP_ALIGN.CENTER)

# ================================================================
# SLIDE 14 — RESULTADOS
# ================================================================
s = new_slide()
bg(s, WHITE)
header(s, "Resultados experimentales", "Validación sobre el prototipo físico")

tests = [
    ("Prueba 1\nEliminación de crosstalk",   "Prueba1.3.png",
     "0 lecturas espurias\nen todos los escenarios"),
    ("Prueba 2\nTolerancia a fallos",         "Prueba2.3.png",
     "Cobertura mantenida\nante pérdida y reconexión"),
    ("Prueba 3\nTrilateración cooperativa",   "Prueba3.4.png",
     "Posición estimada coherente\ncon la geometría del despliegue"),
]

cw = 4.1
for i, (label, im, result) in enumerate(tests):
    x = 0.3 + i * (cw + 0.21)
    rect(s, x, 1.52, cw, 0.65, DARK_BLUE)
    tb(s, label, x + 0.05, 1.57, cw - 0.1, 0.6,
       size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    img(s, im, x, 2.2, w=cw)
    rect(s, x, 5.72, cw, 1.0, LIGHT_BLUE)
    tb(s, result, x + 0.05, 5.77, cw - 0.1, 0.9,
       size=13, bold=True, color=DARK_BLUE, align=PP_ALIGN.CENTER)

# ================================================================
# SLIDE 15 — CONCLUSIONES
# ================================================================
s = new_slide()
bg(s, WHITE)
header(s, "Conclusiones y líneas futuras")

tb(s, "Logros alcanzados", 0.5, 1.52, 6.0, 0.45,
   size=15, bold=True, color=DARK_BLUE)
bullets(s, [
    "El protocolo TDMA dinámico elimina completamente el crosstalk",
    "El sistema tolera desconexiones y reconexiones sin intervención manual",
    "La trilateración cooperativa produce estimaciones coherentes con la geometría real",
    "Todo con hardware de bajo coste (<50 €/nodo) y software de código abierto",
], 0.5, 2.05, 6.0, 4.5, size=15)

rect(s, 6.7, 1.5, 0.06, 5.7, LIGHT_BLUE)

tb(s, "Líneas futuras", 6.9, 1.52, 6.1, 0.45,
   size=15, bold=True, color=DARK_BLUE)
bullets(s, [
    "Ampliar a más de 3 nodos con recálculo dinámico de zonas sucias",
    "Sustituir ultrasonidos por sensores LIDAR o IR para mayor precisión",
    "Clasificación de objetos: tamaño y velocidad estimada",
    "Portar el servidor a un microcontrolador para despliegue autónomo",
], 6.9, 2.05, 6.1, 4.5, size=15)

# ================================================================
# SLIDE 16 — PREGUNTAS
# ================================================================
s = new_slide()
bg(s, DARK_BLUE)
rect(s, 0, 2.85, 13.33, 0.07, MID_BLUE)

tb(s, "¿Preguntas?",
   0.5, 1.1, 12.3, 1.6, size=58, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
tb(s, "Alejandro Millán de Lara",
   0.5, 3.1, 12.3, 0.65, size=20, color=LIGHT_BLUE, align=PP_ALIGN.CENTER)
tb(s, "i12milaa@uco.es",
   0.5, 3.72, 12.3, 0.55, size=16, color=LIGHT_BLUE, align=PP_ALIGN.CENTER)
tb(s, "Sistema de Vigilancia Distribuido Basado en Red de Radares Cooperativos",
   0.5, 4.6, 12.3, 0.7, size=14, color=ACCENT, align=PP_ALIGN.CENTER)

# ================================================================
prs.save(OUTPUT)
print(f"OK -> {OUTPUT}")
